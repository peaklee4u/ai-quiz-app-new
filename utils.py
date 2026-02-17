import PyPDF2
from pptx import Presentation
import io

def extract_text_from_pdf(file):
    """PDF 파일에서 텍스트를 추출합니다."""
    text = ""
    try:
        pdf_reader = PyPDF2.PdfReader(file)
        for page in pdf_reader.pages:
            text += page.extract_text() + "\n"
    except Exception as e:
        return f"Error reading PDF: {e}"
    return text

def extract_text_from_pptx(file):
    """PPTX 파일에서 텍스트를 추출합니다."""
    text = ""
    try:
        prs = Presentation(file)
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"
    except Exception as e:
        return f"Error reading PPTX: {e}"
    return text

def extract_text(uploaded_file):
    """업로드된 파일의 타입에 따라 텍스트를 추출합니다."""
    if uploaded_file.name.endswith('.pdf'):
        return extract_text_from_pdf(uploaded_file)
    elif uploaded_file.name.endswith('.pptx'):
        return extract_text_from_pptx(uploaded_file)
    elif uploaded_file.name.endswith('.txt'):
        return str(uploaded_file.read(), 'utf-8')
    else:
        return "Unsupported file format."
